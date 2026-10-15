import streamlit as st
import pandas as pd
from unidecode import unidecode
import sqlite3
import io
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font
from openpyxl.utils import get_column_letter
from streamlit_autorefresh import st_autorefresh

# ==========================================
# 1. CẤU HÌNH TRANG WEB & TỰ ĐỘNG LÀM MỚI (REAL-TIME)
# ==========================================
st.set_page_config(page_title="Hệ thống Kiểm phiếu tự động", layout="wide")
st.title("📊 Hệ thống Kiểm phiếu Tự động (Đồng bộ đa máy)")

# Tự động quét và làm mới toàn bộ giao diện sau mỗi 3000 mili-giây (3 giây)
st_autorefresh(interval=3000, key="datarefresh")

# ==========================================
# 2. KHỞI TẠO VÀ CẤU HÌNH DATABASE CENTRAL (SQLITE)
# ==========================================
def get_db_connection():
    conn = sqlite3.connect('kiem_phieu_data.db', check_same_thread=False)
    return conn

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS thong_so (
            id INTEGER PRIMARY KEY,
            vong TEXT, phiu_phat INTEGER, phiu_thu INTEGER, 
            phiu_huy INTEGER, phiu_hop_le INTEGER, phiu_trang INTEGER
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS ung_vien (
            ten TEXT PRIMARY KEY,
            so_phieu INTEGER,
            vong TEXT
        )
    ''')
    cursor.execute("SELECT COUNT(*) FROM thong_so")
    res = cursor.fetchone()[0]
    if res == 0:
        cursor.execute("INSERT INTO thong_so VALUES (1, 'Vòng đề cử', 0, 0, 0, 100, 0)")
    conn.commit()
    conn.close()

init_db()

conn = get_db_connection()
df_ts = pd.read_sql_query("SELECT * FROM thong_so WHERE id=1", conn).iloc[0]
danh_sach_uv = pd.read_sql_query("SELECT * FROM ung_vien", conn).to_dict(orient="records")

# ==========================================
# 3. MỤC 1 - THIẾT LẬP THÔNG SỐ & TẢI DANH SÁCH ỨNG VIÊN
# ==========================================
st.header("1. Thiết lập thông số và Tải danh sách ứng viên")

col1, col2, col3, col4, col5, col6 = st.columns(6)
with col1:
    vong_list = ["Vòng đề cử", "Vòng 1", "Vòng 2", "Vòng 3"]
    idx_vong = vong_list.index(df_ts['vong']) if df_ts['vong'] in vong_list else 0
    vong = st.selectbox("Chọn vòng tín nhiệm:", vong_list, index=idx_vong)
with col2:
    phiu_phat_ra = st.number_input("Số phiếu phát ra:", min_value=0, value=int(df_ts['phiu_phat']), step=1)
with col3:
    phiu_thu_vao = st.number_input("Số phiếu thu vào:", min_value=0, value=int(df_ts['phiu_thu']), step=1)
with col4:
    phiu_bat_hop_le = st.number_input("Phiếu bất hợp lệ:", min_value=0, value=int(df_ts['phiu_huy']), step=1)
with col5:
    phiu_hop_le = st.number_input("Số phiếu hợp lệ:", min_value=0, value=int(df_ts['phiu_hop_le']), step=1)
with col6:
    phiu_trang = st.number_input("Số phiếu trắng (hợp lệ):", min_value=0, value=int(df_ts['phiu_trang']), step=1)

tong_mau_so = phiu_hop_le + phiu_trang

col_save_ts, _ = st.columns([2, 10])
with col_save_ts:
    if st.button("💾 Lưu cấu hình thông số", use_container_width=True, type="primary"):
        conn.execute('''
            UPDATE thong_so 
            SET vong=?, phiu_phat=?, phiu_thu=?, phiu_huy=?, phiu_hop_le=?, phiu_trang=? WHERE id=1
        ''', (vong, phiu_phat_ra, phiu_thu_vao, phiu_bat_hop_le, phiu_hop_le, phiu_trang))
        conn.commit()
        st.success("Đã đồng bộ thông số hệ thống!")
        st.rerun()

st.write("")
uploaded_file = st.file_uploader("Tải lên danh sách ứng viên từ Excel (Cột đầu là Họ tên):", type=["xlsx", "xls"])
if uploaded_file and not danh_sach_uv:
    try:
        df = pd.read_excel(uploaded_file)
        ten_ung_vien = df.iloc[:, 0].dropna().tolist()
        for ten in ten_ung_vien:
            conn.execute("INSERT OR IGNORE INTO ung_vien VALUES (?, 0, ?)", (str(ten), vong))
        conn.commit()
        st.success(f"Đã tải thành công {len(ten_ung_vien)} ứng viên vào Database chung!")
        st.rerun()
    except Exception as e:
        st.error(f"Lỗi đọc dữ liệu file Excel: {e}")

col_test, _ = st.columns([2, 10])
with col_test:
    if st.button("✨ Tạo dữ liệu thử nghiệm", use_container_width=True):
        test_names = [f"Nguyễn Văn {chr(65+i)}" for i in range(15)] + ["Nguyễn Văn Tuấn", "Trần Minh Hoàng", "Lê Thị Diễm"]
        for ten in test_names:
            conn.execute("INSERT OR IGNORE INTO ung_vien VALUES (?, 0, ?)", (ten, vong))
        conn.commit()
        st.rerun()

# ==========================================
# 4. MỤC 2 - KIỂM PHIẾU VÀ BẢNG XẾP HẠNG
# ==========================================
if danh_sach_uv:
    st.write("---")
    st.header("2. Kiểm phiếu & Bảng xếp hạng (Dữ liệu chung)")
    
    col_nhap_lieu, col_bxh = st.columns(2)
    
    df_kq = pd.DataFrame(danh_sach_uv)
    df_kq.columns = ["Ten", "SoPhieu", "Vong"]
    
    if tong_mau_so > 0:
        df_kq["Tỷ Lệ %"] = (df_kq["SoPhieu"] * 100 / tong_mau_so).round(2)
    else:
        df_kq["Tỷ Lệ %"] = 0.0
        
    df_kq["Số Phiếu Tín Nhiệm"] = df_kq["SoPhieu"].astype(str) + " / " + str(tong_mau_so)
    df_kq["Đánh Giá"] = df_kq["Tỷ Lệ %"].apply(lambda x: "🟢 ĐẠT" if x >= 50.0 else "🔴 CHƯA ĐẠT")
    
    df_kq = df_kq.sort_values(by="SoPhieu", ascending=False).reset_index(drop=True)
    df_kq.index += 1
    
    df_hien_thi = df_kq[["Ten", "Số Phiếu Tín Nhiệm", "Vong", "Tỷ Lệ %", "Đánh Giá"]]
    df_hien_thi.columns = ["Họ và Tên", "Số Phiếu Tín Nhiệm", "Số Vòng", "Tỷ Lệ %", "Ghi Chú (Đánh giá)"]

    # --- BÊN TRÁI: Ô NHẬP PHIẾU NHANH ---
    with col_nhap_lieu:
        with st.container(border=True):
            st.subheader("⌨️ Ô nhập phiếu nhanh")
            tu_khoa = st.text_input("Gõ tên ứng viên (Không dấu để tìm kiếm):", placeholder="Ví dụ: tuan, hoang, diem...")
            
            if tu_khoa:
                tu_khoa_khong_dau = unidecode(tu_khoa).lower().strip()
                goi_y_khop = [uv for uv in danh_sach_uv if tu_khoa_khong_dau in unidecode(uv["ten"]).lower()]
                
                if goi_y_khop:
                    st.write("*Click chọn đúng tên bên dưới để cộng phiếu:*")
                    for uv_goi_y in goi_y_khop[:5]:
                        ten_goi_y = uv_goi_y["ten"]
                        if st.button(f"👤 {ten_goi_y}", key=f"btn_{ten_goi_y}", use_container_width=True):
                            conn.execute("UPDATE ung_vien SET so_phieu = so_phieu + 1, vong=? WHERE ten=?", (vong, ten_goi_y))
                            conn.commit()
                            st.toast(f"✅ Đã ghi nhận thêm 1 phiếu cho: {ten_goi_y}!", icon="🎉")
                            st.rerun()
                else:
                    st.caption("❌ Không tìm thấy ứng viên nào khớp.")
                
    # --- BÊN PHẢI: BẢNG XẾP HẠNG GỌN GÀNG TRONG KHUNG CÓ THANH CUỘN ---
    with col_bxh:
        with st.container(border=True):
            st.subheader("🏆 Bảng xếp hạng hiển thị")
            st.dataframe(df_hien_thi, use_container_width=True, height=400)

    # ==========================================
    # 5. CHỨC NĂNG XUẤT FILE EXCEL ĐẸP
    # ==========================================
    def export_to_excel_pretty(df_source):
        output = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "KetQuaKiemPhieu"
        
        font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Arial", size=11)
        fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )
        
        headers = ["STT"] + list(df_source.columns)
        ws.append(headers)
        
        for idx, row in enumerate(df_source.itertuples(), start=1):
            ws.append([idx] + list(row[1:]))
            
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.font = font_data
                    if cell.column == 2:
                        cell.alignment = Alignment(vertical="center", horizontal="left")
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    
        # ĐÃ FIX TRIỆT ĐỂ: Duyệt qua số thứ tự cột và dùng get_column_letter chuẩn của openpyxl
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row in range(1, ws.max_row + 1):
                val = ws.cell(row=row, column=col_idx).value
                if val:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        wb.save(output)
        return output.getvalue()

    st.write("")
    excel_data = export_to_excel_pretty(df_hien_thi)
    st.download_button(
        label="📥 Tải về Báo cáo Kết quả (Định dạng Excel có khung kẻ bảng)",
        data=excel_data,
        file_name="Ket_qua_kiem_phieu_chuan.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

# ==========================================
# 6. MỤC 3 - PHẦN CHỈNH SỬA VÀ ĐIỀU CHỈNH DỮ LIỆU
# ==========================================
if danh_sach_uv:
    st.write("---")
    st.header("3. Điều chỉnh dữ liệu (Quản trị viên)")
    
    danh_sach_ten_sua = [uv["ten"] for uv in danh_sach_uv]
    ten_can_sua = st.selectbox("Chọn ứng viên cần sửa số liệu hoặc xóa hẳn khỏi danh sách:", [""] + danh_sach_ten_sua, key="select_sua")
    
    if ten_can_sua:
        info_hien_tai = next(item for item in danh_sach_uv if item["ten"] == ten_can_sua)
        col_info, col_action = st.columns(2)
        
        with col_info:
            st.info(f"Ứng viên đang chọn: **{ten_can_sua}**  \nSố phiếu hiện tại trên hệ thống: **{info_hien_tai['so_phieu']}** phiếu.")
            so_phieu_moi = st.number_input("Nhập chính xác số phiếu thay thế mới:", min_value=0, value=int(info_hien_tai["so_phieu"]), step=1)
            
        with col_action:
            st.write("")
            col_btn_sua, col_btn_xoa = st.columns(2)
            with col_btn_sua:
                if st.button("💾 Lưu thay đổi số phiếu", use_container_width=True, type="primary"):
                    conn.execute("UPDATE ung_vien SET so_phieu=? WHERE ten=?", (so_phieu_moi, ten_can_sua))
                    conn.commit()
                    st.success(f"Đã cập nhật lại số phiếu của {ten_can_sua} thành {so_phieu_moi}!")
                    st.rerun()
            with col_btn_xoa:
                if st.button("🗑️ Xóa ứng viên này", use_container_width=True):
                    conn.execute("DELETE FROM ung_vien WHERE ten=?", (ten_can_sua,))
                    conn.commit()
                    st.warning(f"Đã xóa hoàn toàn ứng viên {ten_can_sua} khỏi danh sách!")
                    st.rerun()

conn.close()
